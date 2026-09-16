"""Builds the Educational Mode copy of a diagnostic pack (xlsx files).

2026-09-16: TallyPrime Educational Mode only saves vouchers dated the 1st, the
2nd or the 31st of a month, and a month without a 31st allows only the 1st and
2nd. The Blossom Retail pack is dated through April, so 87 of its 99 vouchers
fell on days an educational learner could not post. This script re-dates ONLY
the cells under a "Date" header, using the same monotonic rule as
lib/tutor/educational-dates.ts:

    day 1      -> 1
    day 2..16  -> 2
    day 17..31 -> 31 if the month has a 31st, else 2

Monotonic means row order and the bank statement's running balance stay
valid. Every other cell (amounts, balances, narrations, bank references such
as "NEFT/N26040201/...", notes), every style, number format, merged range and
sheet is left exactly as it was. The answer key carries no dates and scoring
never reads the day of month, so the re-dated copy scores against the same
key.

Usage:
    python scripts/build-educational-pack.py <input-folder> <output-folder>

Writes one educational copy per *.xlsx in <input-folder>, then re-reads both
sides and verifies: every rewritten date is on an allowed day, every other
cell (value + style) is identical, sheet/row/column counts match. Exits 1 on
any verification failure. Uploading is a separate, confirmed step:
scripts/upload-educational-pack.mjs.
"""

from __future__ import annotations

import calendar
import re
import sys
from datetime import date, datetime
from pathlib import Path

import openpyxl

MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
TEXT_DATE = re.compile(r"^(\s*)(\d{1,2})-([A-Za-z]{3})-(\d{4})(\s*)$")
# Only used for the report: date-looking fragments left untouched on purpose.
EMBEDDED_DATE = re.compile(r"\b\d{1,2}-[A-Za-z]{3,9}(-\d{2,4})?\b")


def month_has_31(year: int, month: int) -> bool:
    return calendar.monthrange(year, month)[1] == 31


def allowed_days(year: int, month: int) -> set[int]:
    return {1, 2, 31} if month_has_31(year, month) else {1, 2}


def educational_day(year: int, month: int, day: int) -> int:
    if day <= 1:
        return 1
    if day <= 16:
        return 2
    return 31 if month_has_31(year, month) else 2


def map_value(value):
    """Returns (new_value, (year, month, old_day, new_day)) or (value, None) when not a date."""
    if isinstance(value, datetime):
        new_day = educational_day(value.year, value.month, value.day)
        return value.replace(day=new_day), (value.year, value.month, value.day, new_day)
    if isinstance(value, date):
        new_day = educational_day(value.year, value.month, value.day)
        return value.replace(day=new_day), (value.year, value.month, value.day, new_day)
    if isinstance(value, str):
        match = TEXT_DATE.match(value)
        if match and match.group(3).title() in MONTHS:
            lead, day_text, mon, year_text, trail = match.groups()
            year, month, day = int(year_text), MONTHS.index(mon.title()) + 1, int(day_text)
            new_day = educational_day(year, month, day)
            # Keep the source's own padding ("02-Apr-2026" stays two-digit).
            new_day_text = f"{new_day:0{len(day_text)}d}"
            return f"{lead}{new_day_text}-{mon}-{year_text}{trail}", (year, month, day, new_day)
    return value, None


def date_columns(ws):
    """(header_row, column) for every cell whose text is exactly 'Date'."""
    found = []
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.strip().lower() == "date":
                found.append((cell.row, cell.column))
    return found


def rewrite_workbook(src: Path, dst: Path):
    wb = openpyxl.load_workbook(src)
    changes = []  # (sheet, coordinate, old, new)
    skipped = []  # non-empty, non-date cells under a Date header
    date_cells = set()  # every date cell, rewritten or already on an allowed day
    for ws in wb.worksheets:
        for header_row, column in date_columns(ws):
            for row_index in range(header_row + 1, ws.max_row + 1):
                cell = ws.cell(row=row_index, column=column)
                if cell.value is None:
                    continue
                date_cells.add((ws.title, cell.coordinate))
                new_value, info = map_value(cell.value)
                if info is None:
                    skipped.append((ws.title, cell.coordinate, cell.value))
                    continue
                if new_value != cell.value:
                    changes.append((ws.title, cell.coordinate, cell.value, new_value))
                    cell.value = new_value
    dst.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dst)
    return changes, skipped, date_cells


def style_signature(cell):
    return (
        cell.number_format,
        repr(cell.font),
        repr(cell.fill),
        repr(cell.border),
        repr(cell.alignment),
        repr(cell.protection),
        cell.data_type,
    )


def verify(src: Path, dst: Path, expected_changes) -> list[str]:
    problems: list[str] = []
    expected = {(sheet, coord): new for sheet, coord, _old, new in expected_changes}
    original = openpyxl.load_workbook(src)
    rebuilt = openpyxl.load_workbook(dst)
    if original.sheetnames != rebuilt.sheetnames:
        return [f"sheet names differ: {original.sheetnames} vs {rebuilt.sheetnames}"]
    for ws_old in original.worksheets:
        ws_new = rebuilt[ws_old.title]
        if (ws_old.max_row, ws_old.max_column) != (ws_new.max_row, ws_new.max_column):
            problems.append(f"{ws_old.title}: size {ws_old.dimensions} vs {ws_new.dimensions}")
        if set(map(str, ws_old.merged_cells.ranges)) != set(map(str, ws_new.merged_cells.ranges)):
            problems.append(f"{ws_old.title}: merged ranges differ")
        for key, dim in ws_old.column_dimensions.items():
            if ws_new.column_dimensions[key].width != dim.width:
                problems.append(f"{ws_old.title}: column {key} width differs")
        for row_old in ws_old.iter_rows():
            for cell_old in row_old:
                cell_new = ws_new[cell_old.coordinate]
                want = expected.get((ws_old.title, cell_old.coordinate), cell_old.value)
                if cell_new.value != want:
                    problems.append(f"{ws_old.title}!{cell_old.coordinate}: {cell_new.value!r} != {want!r}")
                if style_signature(cell_old) != style_signature(cell_new):
                    problems.append(f"{ws_old.title}!{cell_old.coordinate}: style differs")
        for header_row, column in date_columns(ws_new):
            for row_index in range(header_row + 1, ws_new.max_row + 1):
                value = ws_new.cell(row=row_index, column=column).value
                _same, info = map_value(value)
                if info is None:
                    continue
                year, month, day, _mapped = info
                if day not in allowed_days(year, month):
                    problems.append(f"{ws_new.title}: {value!r} is not on an allowed day")
    return problems


def embedded_dates(path: Path, date_cells: set[tuple[str, str]]):
    wb = openpyxl.load_workbook(path)
    found = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if (ws.title, cell.coordinate) in date_cells or not isinstance(cell.value, str):
                    continue
                if EMBEDDED_DATE.search(cell.value):
                    found.append((ws.title, cell.coordinate, cell.value))
    return found


def main() -> int:
    if len(sys.argv) != 3:
        print(__doc__)
        return 2
    in_dir, out_dir = Path(sys.argv[1]), Path(sys.argv[2])
    sources = sorted(in_dir.glob("*.xlsx"))
    if not sources:
        print(f"No .xlsx files in {in_dir}")
        return 2

    failed = False
    for src in sources:
        dst = out_dir / src.name
        changes, skipped, date_cells = rewrite_workbook(src, dst)
        print(f"\n== {src.name}: {len(changes)} date cell(s) rewritten")
        for sheet, coord, old, new in changes:
            print(f"   {sheet}!{coord}: {old} -> {new}")
        for sheet, coord, value in skipped:
            print(f"   WARNING {sheet}!{coord}: non-date value under a Date header left as is: {value!r}")
        for sheet, coord, value in embedded_dates(src, date_cells):
            print(f"   note {sheet}!{coord}: date inside text left unchanged: {value!r}")
        problems = verify(src, dst, changes)
        if problems:
            failed = True
            print(f"   VERIFY FAILED ({len(problems)}):")
            for problem in problems[:50]:
                print(f"     {problem}")
        else:
            print("   verified: dates on allowed days, all other cells and styles identical")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
